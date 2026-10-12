import io
import zipfile

import docx
import pytest

from loreos import parse_docx
from loreos.parsing import ParseError, Section


def save(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_heading_path_comes_from_the_styles():
    d = docx.Document()
    d.add_paragraph("Before any heading.")
    d.add_heading("Handbook", level=0)  # the Title style
    d.add_paragraph("Handbook intro.")
    d.add_heading("Engineering", level=1)
    d.add_paragraph("Engineering overview.")
    d.add_paragraph("Second paragraph.")
    d.add_heading("On-call", level=2)
    d.add_paragraph("Pages go to the primary.")
    d.add_heading("Escalation", level=3)
    d.add_paragraph("Call the secondary.")
    d.add_heading("Releases", level=2)
    d.add_paragraph("Ship on Tuesdays.")
    d.add_heading("Sales", level=1)
    d.add_paragraph("Sales text.")
    d.add_heading("Ninth", level=9)
    d.add_paragraph("Deepest text.")
    assert parse_docx.parse(save(d)) == [
        Section([], None, "Before any heading."),
        Section(["Handbook"], None, "Handbook intro."),
        Section(["Handbook", "Engineering"], None, "Engineering overview.\n\nSecond paragraph."),
        Section(["Handbook", "Engineering", "On-call"], None, "Pages go to the primary."),
        Section(["Handbook", "Engineering", "On-call", "Escalation"], None, "Call the secondary."),
        Section(["Handbook", "Engineering", "Releases"], None, "Ship on Tuesdays."),
        Section(["Handbook", "Sales"], None, "Sales text."),
        Section(["Handbook", "Sales", "Ninth"], None, "Deepest text."),
    ]


def test_table_rows_are_lines_in_document_order():
    d = docx.Document()
    d.add_heading("Prices", level=1)
    d.add_paragraph("Before the table.")
    table = d.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("Plan", "Price"), ("Team", "10")], strict=False):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    table.cell(2, 0).merge(table.cell(2, 1)).text = "Prices exclude tax."
    d.add_paragraph("After the table.")
    expected = "Before the table.\n\nPlan | Price\nTeam | 10\nPrices exclude tax.\n\nAfter the table."
    assert parse_docx.parse(save(d)) == [Section(["Prices"], None, expected)]


def test_the_cap_is_100_mb():
    assert parse_docx.MAX_UNPACKED_BYTES == 100 * 1024 * 1024


def test_over_the_unpacked_cap_is_refused_before_python_docx_opens_it(monkeypatch):
    d = docx.Document()
    d.add_paragraph("A normal file.")
    data = save(d)
    calls = []

    def opened(*args):
        calls.append(args)
        raise AssertionError("docx.Document was called")

    monkeypatch.setattr(parse_docx, "MAX_UNPACKED_BYTES", 1024)
    monkeypatch.setattr(docx, "Document", opened)
    with pytest.raises(ParseError, match="unpacks to over"):
        parse_docx.parse(data)
    assert calls == []


def zip_without_a_document() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr("notes.txt", "not a Word document")
    return buffer.getvalue()


@pytest.mark.parametrize("data", [b"PK\x03\x04 not really a zip", b"plain text", zip_without_a_document()], ids=["truncated-zip", "not-a-zip", "zip-without-document"])
def test_unreadable_docx_is_a_parse_error(data):
    with pytest.raises(ParseError, match="cannot read the DOCX"):
        parse_docx.parse(data)


def test_no_text_is_a_parse_error():
    d = docx.Document()
    d.add_heading("Only a heading", level=1)
    d.add_paragraph("")
    with pytest.raises(ParseError, match="no text"):
        parse_docx.parse(save(d))
